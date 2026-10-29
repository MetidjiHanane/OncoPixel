"""Word (.docx) research report (python-docx)."""
from __future__ import annotations

import io
from typing import Any, Dict, List, Tuple

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Mm, Pt, RGBColor

from .. import config
from . import report as R
from .strings import tr

MINT = RGBColor(0x0F, 0xA8, 0x9A)
RED = RGBColor(0xB4, 0x23, 0x18)
GREY = RGBColor(0x5B, 0x6B, 0x78)


def to_docx(data: Dict[str, Any]) -> bytes:
    lang = data["language"]
    s = tr(lang)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin = sec.right_margin = Mm(18)
    sec.top_margin, sec.bottom_margin = Mm(18), Mm(18)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10)
    doc.core_properties.title = s["title"]
    doc.core_properties.author = config.APP_NAME

    logo = config.ASSETS_DIR / "logo-light.png"
    if logo.exists():
        doc.add_picture(str(logo), height=Mm(10))
    t = doc.add_paragraph()
    r = t.add_run(s["title"])
    r.bold, r.font.size = True, Pt(18)
    p = doc.add_paragraph()
    r = p.add_run(f"{s['tagline']}  ·  {s['generated']} {data['generated_at'][:16].replace('T', ' ')} UTC  ·  {config.APP_NAME} {config.VERSION}")
    r.font.size, r.font.color.rgb = Pt(8.5), GREY
    if data["test_fixture"]:
        p = doc.add_paragraph()
        r = p.add_run(s["test_model"])
        r.bold, r.font.color.rgb = True, RED

    def heading(text: str) -> None:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        r = p.add_run(text)
        r.bold, r.font.size, r.font.color.rgb = True, Pt(12), MINT

    def table(rows: List[Tuple[str, str]]) -> None:
        tb = doc.add_table(rows=0, cols=2)
        tb.style = "Light List"
        for a, b in rows:
            cells = tb.add_row().cells
            cells[0].text, cells[1].text = a, b
            cells[0].width, cells[1].width = Mm(55), Mm(119)
            for run in cells[0].paragraphs[0].runs:
                run.font.size, run.font.color.rgb = Pt(8.5), GREY

    heading(s["section_patient"]); table(R.patient_rows(data, lang))
    heading(s["section_study"]); table(R.study_rows(data, lang))
    heading(s["section_result"]); table(R.result_lines(data, lang))
    n = doc.add_paragraph(s["confidence_note"])
    n.runs[0].font.size, n.runs[0].font.color.rgb = Pt(8.5), GREY

    imgs = data.get("images", {})
    if imgs:
        heading(s["section_images"])
        order = [("original", s["original"]), ("overlay", s["overlay"]), ("gradcam", s["gradcam"]), ("uncertainty", s["uncertainty"])]
        present = [(k, c) for k, c in order if k in imgs]
        tb = doc.add_table(rows=(len(present) + 1) // 2, cols=2)
        for i, (k, cap) in enumerate(present):
            cell = tb.cell(i // 2, i % 2)
            cell.paragraphs[0].add_run().add_picture(io.BytesIO(imgs[k]), width=Mm(70))
            c = cell.add_paragraph(cap)
            c.runs[0].font.size, c.runs[0].font.color.rgb = Pt(8.5), GREY

    rows = R.measurement_rows(data["measurements"], lang)
    if rows:
        heading(s["section_measure"]); table(rows)
        g = data["segmentation"]
        p = doc.add_paragraph(s["mask_edited"] if g and g["edited"] else s["mask_model"])
        p.runs[0].font.size, p.runs[0].font.color.rgb = Pt(8.5), GREY
    heading(s["section_models"]); table(R.model_rows(data, lang))

    st = data["study"]
    heading(s["section_review"])
    rev = R.clinician_rows(data, lang) + [(s["status"], s["statuses"].get(st["status"], st["status"]))]
    if st["status"] == "reviewed":
        rev += [(s["reviewed_by"], st["reviewed_by"] or "-"), (s["reviewed_at"], (st["reviewed_at"] or "")[:16].replace("T", " ")),
                (s["review_note"], st["review_note"] or "-")]
    else:
        rev.append(("", s["not_reviewed"]))
    table(rev)
    heading(s["section_warnings"])
    for w in R.warning_texts(data, lang) + [s["limits"]]:
        doc.add_paragraph(w, style="List Bullet")

    footer = sec.footer.paragraphs[0]
    footer.text = s["disclaimer"]
    footer.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer.runs[0].font.size, footer.runs[0].font.italic = Pt(7.5), True
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()
